from __future__ import annotations

import re
from functools import lru_cache
from pathlib import Path

from langchain_core.documents import Document

from src.logging_config import get_logger

log = get_logger("qizhi.ingest.loaders")

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".md", ".markdown", ".txt"}

# 国标 PDF（方正书版 Identity-H）常见错误映射特征
_GARBLED_MARKERS = re.compile(r"犌犅|犐犆犛|/G[0-9A-Fa-f]{2,}")
_COMMON_CJK = re.compile(
    r"[的一是了不人我在有他这中大来上国个到说们为子和地出道也时年得就那要下以生会自着去之过家学对可他后能里后心好]"
)
_CJK = re.compile(r"[\u4e00-\u9fff]")


def load_file(
    path: Path,
    *,
    max_pages: int | None = None,
    allow_ocr: bool = True,
) -> list[Document]:
    """按后缀解析单个本地文件为 Document 列表。

    Args:
        max_pages: 仅解析前 N 页（预览用）；None 表示全部。
        allow_ocr: 文本层乱码时是否启用 OCR（入库应 True；快速预览可限页）。
    """
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"不支持的文件类型: {suffix}，支持 {sorted(SUPPORTED_SUFFIXES)}")

    if suffix == ".pdf":
        return _load_pdf(path, max_pages=max_pages, allow_ocr=allow_ocr)
    if suffix == ".docx":
        return _load_docx(path)
    return _load_text(path)


def load_paths(paths: list[Path]) -> list[Document]:
    docs: list[Document] = []
    for p in paths:
        docs.extend(load_file(p))
    return docs


def _looks_garbled(text: str) -> bool:
    """判断 PDF 抽取文本是否为乱码（CID/Identity-H 无 ToUnicode 等）。"""
    raw = (text or "").strip()
    if len(raw) < 8:
        return True

    if _GARBLED_MARKERS.search(raw):
        return True

    compact = re.sub(r"\s+", "", raw)
    if not compact:
        return True

    cjk_chars = _CJK.findall(compact)
    if len(cjk_chars) >= 20:
        common = sum(1 for ch in cjk_chars if _COMMON_CJK.match(ch))
        # 抽出大量「汉字」但几乎不是常用字 → 多为错误码位映射
        if common / max(len(cjk_chars), 1) < 0.08:
            return True

    # 方正书版常见：中文标题被抽成间隔的 ASCII 标点
    spaced_punct = len(re.findall(r"(?:^|\n)[!\"#$%&'()*+,\-./:;<=>?](?:\n|$)", raw))
    if spaced_punct >= 8 and len(cjk_chars) < 30:
        return True

    return False


@lru_cache(maxsize=1)
def _get_ocr():
    from rapidocr_onnxruntime import RapidOCR

    return RapidOCR()


def _ocr_pdf_page(page) -> str:
    """将 PDF 页渲染为图片后 OCR（适用于国标等损坏文本层）。"""
    import pymupdf

    # 2x 分辨率兼顾速度与中文识别率
    pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
    img_bytes = pix.tobytes("png")
    ocr = _get_ocr()
    result, _ = ocr(img_bytes)
    if not result:
        return ""
    return "\n".join(item[1] for item in result if item and item[1]).strip()


def _extract_page_text_pymupdf(page) -> str:
    return (page.get_text("text") or "").strip()


def _load_pdf(
    path: Path,
    *,
    max_pages: int | None = None,
    allow_ocr: bool = True,
) -> list[Document]:
    """优先 PyMuPDF 文本抽取；乱码或空文本时自动 OCR。

    从内存打开 PDF，避免 Windows 下长时间占用 raw 文件导致无法删除。
    """
    try:
        import pymupdf
    except ImportError as e:
        log.warning("pymupdf 不可用，回退 pypdf: %s", e)
        return _load_pdf_pypdf(path, max_pages=max_pages)

    # 读入内存后立即释放路径句柄，避免 WinError 32
    pdf_bytes = path.read_bytes()
    doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    docs: list[Document] = []
    ocr_pages = 0
    try:
        total = doc.page_count
        limit = total if max_pages is None else max(1, min(max_pages, total))
        for i in range(limit):
            page = doc[i]
            text = _extract_page_text_pymupdf(page)
            used_ocr = False
            if allow_ocr and _looks_garbled(text):
                try:
                    ocr_text = _ocr_pdf_page(page)
                    if ocr_text:
                        text = ocr_text
                        used_ocr = True
                        ocr_pages += 1
                except Exception:
                    log.exception("PDF OCR 失败 page=%s file=%s", i + 1, path.name)
            elif (not allow_ocr) and _looks_garbled(text):
                text = ""
            if not text:
                continue
            docs.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": path.name,
                        "source_path": str(path.resolve()),
                        "page": i + 1,
                        "file_type": "pdf",
                        "extract_method": "ocr" if used_ocr else "text",
                        "page_count": total,
                    },
                )
            )
    finally:
        doc.close()

    if ocr_pages:
        log.info(
            "PDF 使用 OCR 解析 pages=%s/%s file=%s",
            ocr_pages,
            len(docs),
            path.name,
        )

    if not docs:
        if not allow_ocr:
            raise ValueError(
                f"PDF 文本层乱码或为空（预览未启用全文 OCR，请先重建索引）: {path.name}"
            )
        raise ValueError(f"PDF 无可提取文本（可能是扫描件且 OCR 失败）: {path}")
    return docs


def _load_pdf_pypdf(path: Path, *, max_pages: int | None = None) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    docs: list[Document] = []
    total = len(reader.pages)
    limit = total if max_pages is None else max(1, min(max_pages, total))
    for i, page in enumerate(reader.pages[:limit], start=1):
        text = (page.extract_text() or "").strip()
        if not text or _looks_garbled(text):
            continue
        docs.append(
            Document(
                page_content=text,
                metadata={
                    "source": path.name,
                    "source_path": str(path.resolve()),
                    "page": i,
                    "file_type": "pdf",
                    "extract_method": "text",
                    "page_count": total,
                },
            )
        )
    if not docs:
        raise ValueError(
            f"PDF 文本层乱码或为空，且当前环境缺少 OCR 依赖（请安装 pymupdf、rapidocr-onnxruntime）: {path}"
        )
    return docs


def _load_docx(path: Path) -> list[Document]:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    text = "\n".join(paragraphs).strip()
    if not text:
        raise ValueError(f"Word 文档无有效文本: {path}")
    return [
        Document(
            page_content=text,
            metadata={
                "source": path.name,
                "source_path": str(path.resolve()),
                "page": 1,
                "file_type": "docx",
            },
        )
    ]


def _load_text(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8").strip()
    if not text:
        raise ValueError(f"空文件: {path}")
    suffix = path.suffix.lower().lstrip(".") or "txt"
    return [
        Document(
            page_content=text,
            metadata={
                "source": path.name,
                "source_path": str(path.resolve()),
                "page": 1,
                "file_type": suffix,
            },
        )
    ]
