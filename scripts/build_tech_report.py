#!/usr/bin/env python
"""生成《智匠·企业知识问答智能体》技术报告 Word 文档。"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "智匠-企业知识问答智能体-技术报告.docx"

BRAND = RGBColor(0x07, 0x4C, 0x37)
INK = RGBColor(0x1B, 0x24, 0x20)
MUTED = RGBColor(0x4A, 0x55, 0x50)
HEADER_FILL = "0B6E4F"


def set_run_font(run, *, name="宋体", size=12, bold=False, color=INK, western="Times New Roman"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = western
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), western)
    rFonts.set(qn("w:hAnsi"), western)
    rFonts.set(qn("w:eastAsia"), name)


def add_para(doc, text, *, size=12, bold=False, first_line=True, space_after=8, color=INK, align="left"):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line:
        pf.first_line_indent = Cm(0.74)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Cm(0)
    if align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pf.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, name="黑体", size=16 if level == 1 else 14, bold=True, color=BRAND, western="Calibri")
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(8)
    h.paragraph_format.first_line_indent = Cm(0)
    return h


def shade_cell(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_run_font(run, name="黑体", size=10.5, bold=True, color=RGBColor(255, 255, 255), western="Calibri")
        shade_cell(cell, HEADER_FILL)
    for r_i, row in enumerate(rows, start=1):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.35
        run = p.add_run(item)
        set_run_font(run, size=12)


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    r = title.add_run("智匠·企业知识问答智能体")
    set_run_font(r, name="黑体", size=22, bold=True, color=BRAND, western="Calibri")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(4)
    r = sub.add_run("技术报告（RAG + Agent）")
    set_run_font(r, name="黑体", size=16, bold=True, color=INK, western="Calibri")

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(16)
    r = meta.add_run("面向企业内部制度问答  ·  可演示、可评测、可离线部署")
    set_run_font(r, size=10.5, color=MUTED)

    add_heading(doc, "一、面试介绍（开场 60 秒）", 1)
    add_para(
        doc,
        "面试官您好。我做的这个项目叫「智匠」，是一套面向企业内部制度、规范和操作手册的知识问答智能体，"
        "核心不是把文档丢进大模型，而是用 RAG 解决依据从哪来、用 Agent 解决怎么检索、用工程约束解决回答能不能追溯。"
        "员工用自然语言提问后，Agent 会按问法选择工具：口语简称或含指代时先 rewrite_query 改写成检索词，"
        "再 search_knowledge 做 BM25 与向量的混合召回；如果只是问「库里有哪些制度」，则走 list_documents。"
        "生成阶段对接内网 vLLM，系统在后端把「来源：知识库/文件名」拼到回答末尾，不让模型自己编引用。"
        "产品侧有 JWT 角色鉴权、多知识库隔离、流式工具时间线、Admin 召回对比，以及 Docker 单镜像离线部署。"
        "检索效果我在本机用 32 道黄金集实跑过对照（当时未加载 reranker）：相对纯向量，BM25+向量混合检索 Recall@1 从 75.0%（24/32）提到 81.2%（26/32），Recall@5 从 96.9%（31/32）提到 100%。"
        "这是小规模构造集，面试时我会说明口径。如果现场要看，可以演示「先改写再检索」的时间线，也可以打开一条失败轨迹：问「和去年比有没有变化」时系统拒答，而不是编造条款。",
    )

    add_heading(doc, "二、项目定位与要解决的问题", 1)
    add_para(
        doc,
        "企业内部制度通常散落在 PDF、Word、Markdown 中。员工检索成本高，直接把全文塞给大模型又会超窗、不可审计。"
        "本项目把「能问、有据、可管、可部署」拆成四件事：",
    )
    add_bullets(doc, [
        "能问：自然语言多轮问答，而不是关键词搜索。",
        "有据：回答必须绑定检索到的制度切片，来源可追溯；知识库没有的内容明确拒答。",
        "可管：管理员维护多个知识库，普通员工只能问自己会话，不能改索引。",
        "可部署：对话模型走内网 vLLM，本机只跑 Embedding 与业务服务；支持无网机 Docker 加载。",
    ])
    add_para(
        doc,
        "当前版本为 V2。相对早期「全员一份知识库、无登录」的形态，已补齐 JWT 鉴权、会话按用户隔离、多知识库分集合，以及混合检索、流式工具轨迹和黄金集评测。",
    )

    add_heading(doc, "三、总体架构", 1)
    add_para(
        doc,
        "系统分表现层、接入层、智能体层、检索层、入库层和身份层。端到端链路如下：",
        first_line=True,
    )
    add_para(
        doc,
        "管理员上传文档 → 原文落入 data/kbs/<知识库ID>/raw/（此时不做 Embedding）→ 点击重建索引："
        "解析 / 中文分块 / BGE 向量化 / 写入 Chroma 对应集合。员工登录后发起会话并绑定知识库 →"
        "提问进入 LangChain Agent → 按需调用 rewrite_query / search_knowledge / list_documents →"
        "混合检索返回切片 → 内网 LLM 生成正文 → 后端按工具结果拼接来源 → LangGraph Checkpoint 与会话表持久化。",
        first_line=True,
    )
    add_para(doc, "逻辑分层与职责：", first_line=True)
    add_table(
        doc,
        ["层次", "职责", "主要实现"],
        [
            ["表现层", "问答、知识库管理、登录注册、系统管理", "Next.js 15 + TypeScript + Ant Design 6"],
            ["接入层", "REST API、流式 NDJSON、鉴权依赖、上传预览", "FastAPI + Uvicorn，默认 :8002"],
            ["身份层", "用户、JWT、角色、审计日志", "PyJWT + bcrypt + SQLite users/audit_logs"],
            ["智能体层", "工具选择、多轮记忆、来源后处理、流式事件", "LangChain create_agent + LangGraph SqliteSaver"],
            ["检索层", "词法召回、向量召回、融合与可选重排", "jieba + rank-BM25 + Chroma + BGE / 可选 reranker"],
            ["入库层", "PDF/Word/MD 解析、乱码检测、OCR、分块", "pypdf / PyMuPDF / python-docx / RapidOCR"],
            ["基础设施", "配置、vLLM 客户端、日志、本地轨迹", ".env、OpenAI 兼容接口、滚动日志、data/traces"],
        ],
    )
    add_para(
        doc,
        "两个刻意拆开的边界：上传不等于向量化，避免大文件阻塞问答；大模型不在本机，本机 CPU 只承担 Embedding 与业务逻辑，便于内网合规部署。",
    )

    add_heading(doc, "四、技术选型与使用方式", 1)

    add_heading(doc, "4.1 技术栈一览", 2)
    add_table(
        doc,
        ["类别", "技术", "在本项目中的用法"],
        [
            ["对话模型", "vLLM（OpenAI 兼容），默认 Qwen3.5-VL-122B", "Chat Completions + 工具调用；需开启 tool-call-parser"],
            ["Agent 框架", "LangChain 1.x create_agent", "注册三个工具，由模型决定调用顺序"],
            ["状态机/记忆", "LangGraph SqliteSaver", "按 thread_id 持久化完整 message 轨迹，支持多轮"],
            ["词法检索", "jieba + rank-BM25", "中文切词后建内存倒排，补精确条款、数字、专名"],
            ["向量检索", "BAAI/bge-small-zh-v1.5 + Chroma", "本地 CPU Embedding，按知识库分 collection 持久化"],
            ["融合 / 重排", "RRF；可选 bge-reranker-base", "两路排名融合后再交叉编码（权重缺失则自动降级）"],
            ["后端", "Python 3.11+、FastAPI、Pydantic", "同步问答 + chat/stream 流式接口"],
            ["前端", "Next.js 15 App Router、React 19", "rewrite 代理到 API；问答页渲染工具时间线"],
            ["鉴权", "JWT（HS256）+ bcrypt", "access token 与 kb_admin token 分类型校验"],
            ["存储", "文件系统 + Chroma + SQLite", "原文、向量、用户/审计/会话/Checkpoint 同机持久化"],
            ["部署", "Docker 单镜像（API+Web+Embedding）", "vLLM 仍走内网；数据卷升级不丢库"],
            ["评测与观测", "自研黄金集脚本、pytest、Langfuse 可选", "Recall@K、失败轨迹、鉴权/来源单测"],
        ],
    )

    add_heading(doc, "4.2 为什么用 Agent，而不是一条 RAG 链", 2)
    add_para(
        doc,
        "如果只做 retrieve-then-generate，用户原话会被直接 embedding。制度场景里大量问法是「年假几天」「这个和去年比呢」——简称、指代、对比。"
        "Agent 的价值是路由：先判断要不要改写、要不要列目录、要不要检索，而不是每次都把原句塞进同一个向量查询。"
        "三个工具的职责如下。",
    )
    add_table(
        doc,
        ["工具", "触发场景", "输入 / 输出"],
        [
            ["rewrite_query", "简称、口语、指代、对比问法", "输入原问题；输出主查询、备选查询、是否需要结合上文"],
            ["search_knowledge", "制度/流程/操作规范类问题", "用改写后的查询做混合检索，返回带「来源/页段」的切片"],
            ["list_documents", "「有哪些制度」「库里有什么文件」", "列出当前知识库已入库文档名"],
        ],
    )
    add_para(
        doc,
        "改写本身不额外调用大模型，规则可单测：例如「年假」展开为「年休假」，「和去年比呢」会标记 needs_context 并补上「结合上文主题 / 制度条款是否变更」。"
        "这样既降低一次 LLM 往返，也让面试时能讲清 Agent 比链式 RAG 多了哪一步决策。",
    )

    add_heading(doc, "4.3 混合检索怎么做", 2)
    add_para(
        doc,
        "默认链路：对同一查询分别取向量 Top-N 与 BM25 Top-N（默认候选 20），用 Reciprocal Rank Fusion（RRF，k=60）融合；"
        "若本地存在 models/bge-reranker-base，再对融合结果做 CrossEncoder 重排，截断到 RETRIEVE_TOP_K（默认 4）交给 Agent。"
        "没有 reranker 时自动降级为混合检索，不阻断启动。Admin 召回页会并排展示「仅向量 / BM25+向量 / 混合+Rerank」，便于对照失败 case。",
    )

    add_heading(doc, "4.4 来源为什么不能交给模型写", 2)
    add_para(
        doc,
        "制度问答最怕「条款是对的、出处是编的」。prompt 明确禁止模型在文末自行添加「来源」行；"
        "后端从 search_knowledge 的 ToolMessage 与检索 metadata 解析文件名、知识库名、页/段，在 finalize_answer_sources 里剥掉模型误写的来源后再追加真实标签，格式为「来源：{知识库}/{文件名}」。",
    )

    add_heading(doc, "4.5 大模型接入要点", 2)
    add_para(
        doc,
        "llm.py 使用 LangChain init_chat_model，model_provider=openai，base_url 指向内网 vLLM。"
        "Agent 依赖服务端工具调用能力，部署时需要 --enable-auto-tool-choice 以及对应 Qwen 的 --tool-call-parser qwen3_xml（或等价解析器）。"
        "这是 Agent 能否真正「先检索再作答」的基础设施条件，不是 prompt 能单独解决的。",
    )

    add_heading(doc, "五、已实现功能", 1)

    add_heading(doc, "5.1 面向员工的问答", 2)
    add_bullets(doc, [
        "登录后多轮对话；每个会话绑定一个知识库，历史按用户隔离。",
        "POST /api/v1/chat/stream 以 NDJSON 推送 status / tool_start / tool_end / token / done；前端展示「改写检索词 → 检索知识库 → 生成回答」时间线。",
        "同步接口 /api/v1/chat 保留，便于脚本调用。",
        "身份类问题（你是谁、什么模型）按系统提示直接回答，不调用工具，且不暴露底层模型品牌。",
    ])

    add_heading(doc, "5.2 知识库管理", 2)
    add_bullets(doc, [
        "支持多个知识库：独立名称、独立原文目录、独立 Chroma collection。",
        "上传与重建索引解耦；入库进度以流式 NDJSON 回传前端。",
        "原文预览、切片预览、按文件删除原文与向量。",
        "召回测试：同一问题对比多路检索结果及分数类型（距离 / RRF / Rerank）。",
    ])

    add_heading(doc, "5.3 文档解析", 2)
    add_bullets(doc, [
        "PDF：优先文本层；检测方正书版 / Identity-H 乱码特征后走 RapidOCR。",
        "Word（.docx）、Markdown、纯文本。",
        "RecursiveCharacterTextSplitter，中文标点分隔，默认 chunk 700 / overlap 100。",
    ])

    add_heading(doc, "5.4 身份认证与权限", 2)
    add_table(
        doc,
        ["能力", "普通用户", "管理员"],
        [
            ["注册 / 登录问答", "可以（只能注册 user）", "可以"],
            ["自己的会话 CRUD、流式问答", "可以", "可以"],
            ["知识库上传 / 重建 / 召回测试", "需管理员二次授权", "可以直接进入"],
            ["用户管理、审计日志", "不可以", "可以"],
            ["自助注册为管理员", "不可以", "管理员由环境变量初始化"],
        ],
    )
    add_para(
        doc,
        "安全边界在 FastAPI Depends，而不是前端路由。Access Token 与知识库临时 Token 类型字段不同，"
        "kb_admin token 不能当作普通登录凭证使用。密码 bcrypt 存储；登录、授权、用户变更写入 audit_logs。",
    )

    add_heading(doc, "5.5 检索黄金集评测（本机实测）", 2)
    add_para(
        doc,
        "下列数字来自本机执行 python scripts/eval_retrieval.py，原始 JSON 见 data/eval/last_report.json，"
        "不是估算值。当时 reranker_ready=false，本机没有 bge-reranker-base，因此对照的是「仅向量」与「BM25+向量 RRF」，"
        "不能写成「重排后提升」。",
    )
    add_para(doc, "评测怎么构造：", first_line=True)
    add_bullets(doc, [
        "题目：evals/golden_set.json 共 32 题，覆盖休假、报销、入职三类示例制度。",
        "语料：samples/ 三篇制度 + evals/distractors/ 五篇干扰文档（外包考勤、子公司差旅、实习生说明等，数字和条款故意写反）。",
        "分块：评测入库用 chunk_size=160，避免示例文档太短、每篇只有 1 个切片导致 Recall 恒为 100%。",
        "隔离：写入独立 Chroma 集合 eval_golden，不污染业务知识库。",
        "指标：Recall@K = 期望文件名是否出现在 Top-K；Support@5 = must_contain 关键词是否出现在 Top-5 切片中（检索侧忠实度代理，不是生成侧 RAGAS）。",
    ])
    add_table(
        doc,
        ["模式", "Recall@1", "Recall@3", "Recall@5", "Support@5", "漏题"],
        [
            ["仅向量 similarity_search", "75.0%（24/32）", "96.9%（31/32）", "96.9%（31/32）", "93.8%", "onboard-05"],
            ["BM25 + 向量 RRF（hybrid）", "81.2%（26/32）", "100%（32/32）", "100%（32/32）", "100%", "无"],
            ["混合 + bge-reranker", "未测", "未测", "未测", "未测", "权重未加载"],
        ],
    )
    add_para(
        doc,
        "如何解读：Recall@5 的提升其实只多对了 1 题（31/32→32/32），说服力有限；更站得住的是 Recall@1 从 24/32 提到 26/32。"
        "这是小规模构造集，用来验证混合检索对「简称/精确数字/干扰文档」的帮助，不能外推成生产制度库的效果。"
        "复现命令：在 enterprise_kb_agent 目录执行 python scripts/eval_retrieval.py --k 5。",
    )

    add_heading(doc, "5.6 单测与失败轨迹", 2)
    add_para(
        doc,
        "工程侧用 pytest 锁住鉴权与来源约束，避免「演示能跑、回归全无」。本机 python -m pytest tests -q 结果为 16 passed。",
    )
    add_table(
        doc,
        ["文件", "覆盖点", "断言意图"],
        [
            ["tests/test_auth.py", "bcrypt 哈希往返、JWT 签发/解码、过期 Token、冻结账号、kb_admin 不能当登录凭证、非管理员 403", "权限漏洞不能靠前端隐藏"],
            ["tests/test_sources.py", "工具输出解析来源、模型自编「来源：」被剥离、真实来源追加、同文件切片去重", "引用必须绑定检索 metadata"],
            ["tests/test_rewrite_query.py", "「年假」展开为年休假、指代问法标记 needs_context、空输入不崩溃", "改写规则可单测、不依赖 LLM"],
        ],
    )
    add_para(
        doc,
        "失败轨迹样例（evals/sample_fail_trace.json，可由 python scripts/demo_fail_trace.py 重生）："
        "用户问「这个和去年比有变化吗？」→ rewrite_query 补上「结合上文主题 / 制度条款是否变更」→ search_knowledge 返回空 → 输出「知识库中未找到相关规定」。"
        "用于说明：知识库没有同比变更条款时必须拒答。配置 LANGFUSE_PUBLIC_KEY / SECRET_KEY 后可同步到 Langfuse；未配置只写 data/traces/，不影响启动。",
    )

    add_heading(doc, "5.7 部署形态", 2)
    add_bullets(doc, [
        "开发：start.bat / start.ps1 / start.sh 同终端拉起 API 与 Web，健康检查通过后再启前端。",
        "生产/内网无网机：单 Docker 镜像内同时跑 FastAPI(:8002) 与 Next.js(:3000)，Embedding 打进镜像，vLLM 走环境变量。",
        "数据卷持久化原文、Chroma、SQLite，镜像升级不丢知识库。",
    ])

    add_heading(doc, "六、关键实现细节（面试可展开）", 1)
    add_heading(doc, "6.1 流式协议", 2)
    add_para(
        doc,
        "chat/stream 使用 application/x-ndjson，事件类型包括 status、tool_start、tool_end、token、done、error。"
        "工具名在前端映射为中文：rewrite_query→改写检索词，search_knowledge→检索知识库，list_documents→列出文档。"
        "最终 done.answer 会带上后处理过的来源行，避免流式过程中引用漂移。",
    )
    add_heading(doc, "6.2 知识库隔离方式", 2)
    add_para(
        doc,
        "每个知识库有 kb_id、显示名、collection_name。检索前根据会话上的 kb_id 写入 ContextVar，"
        "search_knowledge 读当前 collection，避免多库串味。会话表记录 user_id 与 kb_id，接口校验线程归属。",
    )
    add_heading(doc, "6.3 失败怎么讲", 2)
    add_para(
        doc,
        "推荐准备两个故事。其一，国标 PDF 抽出来是乱码：先用规则判断 Identity-H / 非常用汉字比例，再 OCR，而不是全量 OCR。"
        "其二，用户问同比变化但制度里没有该条款：Agent 改写后仍召回空，必须拒答。这两条都能说明系统在「效果」和「诚实」之间做了取舍。",
    )

    add_heading(doc, "七、主要接口（便于联调）", 1)
    add_table(
        doc,
        ["方法", "路径", "说明"],
        [
            ["GET", "/health", "健康检查，无需登录"],
            ["POST", "/api/v1/auth/register | /login | /me", "注册、登录、当前用户"],
            ["CRUD", "/api/v1/threads", "会话列表/创建/删除，按登录用户隔离"],
            ["POST", "/api/v1/chat", "同步问答"],
            ["POST", "/api/v1/chat/stream", "流式问答（工具时间线 + token）"],
            ["CRUD", "/api/v1/admin/kbs", "多知识库管理（管理员）"],
            ["GET", "/api/v1/admin/kbs/{id}/recall", "召回对比"],
            ["POST", "/api/v1/admin/kbs/{id}/upload | /ingest", "上传与重建索引"],
            ["GET", "/api/v1/admin/users | /audit-logs", "用户与审计"],
        ],
    )

    add_heading(doc, "八、总结", 1)
    add_para(
        doc,
        "本项目把企业知识问答做成完整产品：入库可运维、检索可对比、Agent 可路由、回答可追溯、权限可隔离、部署可离线。"
        "技术上的重点不是堆框架名称，而是三条约束——问法先改写再检索、两路召回再融合、来源由系统绑定而不是模型生成。"
        "后续若继续迭代，优先补齐 reranker 权重后重跑黄金集第三列、对话生成侧的忠实度打分，以及按部门授权可见知识库；"
        "不打算为了简历去加多 Agent Supervisor 或把 MCP 硬接到本项目上。",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"[OK] {path}")
